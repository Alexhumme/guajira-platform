# -*- coding: utf-8 -*-
"""Genera el documento de entrega de la plataforma Comured."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ---------- utilidades de formato ----------
def set_cell_bg(cell, hex_color):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_bg(cell, '1F4E79')
    for r in rows:
        row = t.add_row().cells
        for i, v in enumerate(r):
            row[i].text = str(v)
    if widths:
        for i, w in enumerate(widths):
            for r in t.rows:
                r.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t

def h(text, level=1):
    doc.add_heading(text, level=level)

def p(text=''):
    doc.add_paragraph(text)

def bullets(items):
    for it in items:
        doc.add_paragraph(it, style='List Bullet')

# ---------- portada ----------
title = doc.add_heading('DOCUMENTO DE ENTREGA TÉCNICA', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph('Plataforma Comured')
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].italic = True
doc.add_paragraph()
table(['Campo', 'Detalle'], [
    ['Proyecto', 'IAP — Investigación Acción Participativa, Zona Caribe 2'],
    ['Entidad', 'SENA — Área SENAINNOVA'],
    ['Rol', 'Técnico en Programación de Software'],
    ['Desarrollador', 'Alex Valdelamar Bustamante (contratista)'],
    ['Fecha de entrega', 'Octubre de 2026'],
    ['Aplicación', 'Plataforma web Comured (api + web + panel admin)'],
])

# ---------- 1 ----------
h('1. Descripción general', 1)
p('Comured es una plataforma digital que visibiliza y gestiona las comunidades rurales de La Guajira: '
  'presentación de comunidades y sus municipios, marketplace de productos artesanales y tradicionales, '
  'rutas turísticas, publicaciones comunitarias, asociaciones y monitoreo de madurez comunitaria. '
  'Incluye un sitio público (landing), una API REST y un panel de administración para cargar y gestionar la información.')

# ---------- 2 ----------
h('2. Arquitectura del sistema', 1)
table(['Componente', 'Tecnología', 'Descripción'], [
    ['api/', 'Node.js + Express 5 + MySQL2', 'API REST (puerto 5000). Sirve endpoints /api/* (admin), /web-client/* (público), /uploads/* (assets).'],
    ['web/', 'Next.js 16 + React 19 + Tailwind 4', 'Sitio público. Se compila como export estático (web/out).'],
    ['admin-react/', 'Vite + React 19', 'Panel de administración (CRUDs, login con sesión MySQL).'],
    ['app/', 'React Native 0.82', 'Aplicación móvil (en estado inicial, sin integración aún).'],
    ['MySQL', 'Base de datos comured', '20 tablas: admin, rol, tipo_producto, departamento, municipio, comunidad, miembro, producto, post, ruta, asociacion, etc.'],
])

# ---------- 3 ----------
h('3. Recursos y direcciones (URLs)', 1)
table(['Recurso', 'URL', 'Notas'], [
    ['Landing (sitio público)', 'https://comured.appsennovaguajira.com', 'Alojada en Hostinger (plan Cloud Startup), contenido estático.'],
    ['API (backend)', 'https://api.<dominio-de-produccion>', 'Node.js en Hostinger Cloud Startup (puerto 5000 interno).'],
    ['Panel de administración', 'https://<dominio>/admin o build local (npm run dev)', 'Vite build → public/admin/dist.'],
    ['Base de datos', 'MySQL en el hosting (localhost:3306), BD: comured', 'Gestión vía phpMyAdmin de Hostinger.'],
    ['Repositorio de código', 'https://github.com/Alexhumme/guajira-platform', 'Rama main.'],
    ['Documentación técnica (diagramas Mermaid)', 'documentacion/', 'Casos de uso, tablas, flujos, vistas y formularios.'],
])

# ---------- 4 ----------
h('4. Credenciales de acceso', 1)
p('Importante: rotar las contraseñas por defecto antes de dejar el sistema en producción.')
table(['Sistema', 'Usuario / clave', 'Detalle'], [
    ['Panel admin (login)', 'Usuario: admin / Clave: admin@1324', 'Definido por ADMIN_BOOTSTRAP_USER / ADMIN_BOOTSTRAP_PASSWORD en api/.env'],
    ['Base de datos MySQL', 'root / jbldb@0173A+', 'Usuario local del servidor MySQL (hosting).'],
    ['Sesión', 'SESSION_SECRET configurable', 'En api/.env (SESSION_SECRET).'],
    ['phpMyAdmin (Hostinger)', 'Credenciales del panel de Hostinger', 'Acceso desde hPanel → Bases de datos.'],
])

# ---------- 5 ----------
h('5. Base de datos', 1)
p('Base de datos: comured (charset utf8mb4). Para importarla en un entorno nuevo:')
bullets([
    'CREATE DATABASE comured CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;',
    'Importar api/sql/schema.sql (estructura) y api/sql/seed.sql (datos de prueba) vía phpMyAdmin o consola.',
    'Las credenciales de conexión están en api/.env (DB_HOST, DB_USER, DB_PASSWORD, DB_NAME, DB_PORT).',
])

# ---------- 6 ----------
h('6. Despliegue y entrega en Hostinger', 1)
bullets([
    'Frontend: con NEXT_PUBLIC_API_URL=<URL pública de la API> en web/.env, ejecutar npm run build en web/ y subir el contenido de web/out/ al document root del subdominio.',
    'API: desplegar la app Node api/ en Hostinger Cloud Startup (Node 20.x), configurar las variables de entorno equivalentes a api/.env y arrancar con node server.js (o el gestor de apps del hosting).',
    'Base de datos: subir el script api/sql/schema.sql y seed.sql a phpMyAdmin del hosting (misma cuenta Cloud Startup).',
    'Panel admin: npm run build en admin-react/ y servir el dist como sitio estático (o integrar en el mismo dominio).',
])

# ---------- 7 ----------
h('7. Explicación de uso de la plataforma', 1)
h('7.1 Visitante (sitio público)', 2)
bullets([
    'Explorar comunidades en /comunidades con filtros por municipio/asociación.',
    'Ver el detalle de cada comunidad (líderes, galería, rutas, productos).',
    'Consultar el marketplace de productos en /marketplace.',
    'Revisar publicaciones en /publicaciones y conocer el proyecto en /proyecto.',
])
h('7.2 Administrador (panel)', 2)
bullets([
    'Iniciar sesión con las credenciales de la sección 4.',
    'Gestionar CRUDs: comunidades, miembros, productos, rutas, categorías, municipios, roles, asociaciones y posts.',
    'Subir logotipos/portadas/galerías mediante los campos "media" de cada formulario.',
    'Monitorear la madurez de las comunidades y exportar datos a Excel.',
])
h('7.3 Miembro / líder comunitario', 2)
bullets([
    'Acceso mediante usuario registrado por el administrador.',
    'Publicar productos y posts asociados a su comunidad.',
])

# ---------- 8 ----------
h('8. Estructura del repositorio', 1)
table(['Ruta', 'Contenido'], [
    ['api/server.js', 'Entrada del servidor Express.'],
    ['api/routes/', 'Rutas /web-client (públicas) y /api (admin).'],
    ['api/sql/', 'schema.sql y seeders.'],
    ['web/app/(public)/', 'Páginas del sitio público.'],
    ['web/lib/api/', 'Clientes de datos (fetch + safeFetchApi).'],
    ['web/hooks/', 'Hooks de React para datos en runtime.'],
    ['admin-react/src/', 'Panel de administración (componentes y secciones).'],
    ['documentacion/', 'Diagramas Mermaid y documentación de formularios.'],
    ['README.md', 'Guía de instalación y puesta en marcha.'],
])

# ---------- 9 ----------
h('9. Mantenimiento y soporte', 1)
bullets([
    'Actualizar dependencias con npm audit y pruebas después de cada cambio.',
    'Rotar SESSION_SECRET y credenciales periódicamente.',
    'Respaldos: exportar la BD desde phpMyAdmin o mysqldump de forma semanal.',
    'Logs: el servidor registra peticiones con morgan; revisar en el panel de Hostinger.',
])

doc.save('ENTREGA_Comured.docx')
print('ENTREGA_Comured.docx generado')
